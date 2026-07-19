"""报价单导出服务 — 生成 Excel (.xlsx) 报价单。"""

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.generation_run import GenerationRun
from app.models.project import Project


# ── 样式常量 ──
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

TITLE_FONT = Font(name="微软雅黑", size=16, bold=True)
HEADER_FONT = Font(name="微软雅黑", size=10, bold=True)
BODY_FONT = Font(name="微软雅黑", size=10)
TOTAL_FONT = Font(name="微软雅黑", size=10, bold=True)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT_WHITE = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
LABEL_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")


def _set_cell(ws, row: int, col: int, value, font=BODY_FONT, alignment=CENTER, border=THIN_BORDER, number_format=None):
    """Helper to set cell value and style."""
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font
    cell.alignment = alignment
    cell.border = border
    if number_format:
        cell.number_format = number_format
    return cell


class QuoteExportService:
    """生成 Excel 报价单。"""

    def __init__(self, db: Session) -> None:
        self.db = db

    def export_selected_plan(self, project_id: str) -> bytes:
        """导出项目已选方案的报价单为 Excel 字节流。"""
        project = self.db.query(Project).filter(Project.id == project_id).first()
        if not project:
            raise ValueError("项目不存在")
        if project.stage != "completed" or not project.selected_run_id or not project.selected_scenario_id:
            raise ValueError("项目尚未完成方案确认，无法导出")

        # 加载定价运行
        run = (
            self.db.query(GenerationRun)
            .filter(
                GenerationRun.id == project.selected_run_id,
                GenerationRun.project_id == project_id,
                GenerationRun.status == "succeeded",
            )
            .first()
        )
        if not run or not run.pricing_payload:
            raise ValueError("定价运行记录不存在或数据缺失")

        # 从 payload 中找到已选方案
        plans = run.pricing_payload.get("plans", [])
        plan = next((p for p in plans if p.get("id") == project.selected_scenario_id), None)
        if not plan:
            raise ValueError("未找到已选方案")

        return self._generate_workbook(project, plan)

    def export_plan_from_payload(
        self, project: Project, plan: dict, pricing_payload: dict
    ) -> bytes:
        """直接从已有数据生成 Excel（用于导出未选定的方案）。"""
        return self._generate_workbook(project, plan)

    def _generate_workbook(self, project: Project, plan: dict) -> bytes:
        """生成 Excel 工作簿。"""
        wb = Workbook()
        ws = wb.active
        ws.title = "报价单"

        # ── 列宽 ──
        col_widths = {1: 6, 2: 40, 3: 12, 4: 16, 5: 12, 6: 12, 7: 14}
        for col, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col)].width = width

        row = 1
        # ── 标题 ──
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        title_cell = ws.cell(row=row, column=1, value="报 价 单")
        title_cell.font = TITLE_FONT
        title_cell.alignment = CENTER
        row += 2

        # ── 项目信息 ──
        info_data = [
            ("报价单位", project.quote_company),
            ("报价日期", str(project.quote_date) if isinstance(project.quote_date, date) else str(project.quote_date)),
        ]
        if project.name:
            info_data.insert(0, ("项目名称", project.name))
        if project.customer_name:
            info_data.append(("客户名称", project.customer_name))

        for label, value in info_data:
            _set_cell(ws, row, 1, label, font=HEADER_FONT, alignment=LEFT, border=THIN_BORDER)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
            _set_cell(ws, row, 2, value, font=BODY_FONT, alignment=LEFT, border=THIN_BORDER)
            row += 1
        row += 1

        # ── 价格对比 ──
        target_cents = project.target_gross_cents
        gross_cents = plan.get("gross_cents", 0)
        labor_cents = plan.get("labor_cents", 0)
        tax_cents = plan.get("tax_cents", 0)
        within_target = plan.get("within_target", False)
        gap = abs(gross_cents - target_cents)

        summary_data = [
            ("目标报价", f"{(target_cents / 100 / 10000):.2f} 万元"),
            ("方案报价", f"{(gross_cents / 100 / 10000):.2f} 万元"),
            ("差额", f"{(gap / 100 / 10000):.2f} 万元　【{'√ 达标' if within_target else '未达标'}】"),
        ]
        for label, value in summary_data:
            _set_cell(ws, row, 1, label, font=HEADER_FONT, alignment=LEFT, border=THIN_BORDER)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
            _set_cell(ws, row, 2, value, font=BODY_FONT, alignment=LEFT, border=THIN_BORDER)
            row += 1
        row += 1

        # ── 明细表头 ──
        headers = ["序号", "功能模块", "角色", "单价（元/人天）", "半天数", "折合天数", "小计（元）"]
        for ci, h in enumerate(headers, 1):
            _set_cell(ws, row, ci, h, font=HEADER_FONT_WHITE, alignment=CENTER, border=THIN_BORDER, number_format="@")
            ws.cell(row=row, column=ci).fill = HEADER_FILL
        header_row = row
        row += 1

        # ── 明细行 ──
        lines = plan.get("lines", [])
        for idx, ln in enumerate(lines, 1):
            wp_name = ln.get("work_package_name") or ln.get("work_package_id") or ""
            role = ln.get("role", "")
            price = ln.get("unit_price_cents", 0)
            half_days = ln.get("half_day_units", 0)
            days = half_days / 2
            subtotal = round(price * half_days / 2)

            _set_cell(ws, row, 1, idx, font=BODY_FONT, alignment=CENTER)
            _set_cell(ws, row, 2, wp_name, font=BODY_FONT, alignment=LEFT)
            _set_cell(ws, row, 3, role, font=BODY_FONT, alignment=CENTER)
            _set_cell(ws, row, 4, f"{price / 100:.0f}", font=BODY_FONT, alignment=CENTER)
            _set_cell(ws, row, 5, half_days, font=BODY_FONT, alignment=CENTER)
            _set_cell(ws, row, 6, f"{days:.1f}", font=BODY_FONT, alignment=CENTER)
            _set_cell(ws, row, 7, f"{subtotal / 100:.2f}", font=BODY_FONT, alignment=RIGHT)
            row += 1

        # ── 汇总行 ──
        data_start = row
        for _ in range(2):  # 留空行
            row += 1

        _set_cell(ws, row, 1, "汇总", font=TOTAL_FONT, alignment=CENTER)
        for ci in range(2, 6):
            _set_cell(ws, row, ci, "", font=TOTAL_FONT, alignment=CENTER)
        _set_cell(ws, row, 6, "人工费", font=TOTAL_FONT, alignment=CENTER)
        _set_cell(ws, row, 7, f"{labor_cents / 100:.2f}", font=TOTAL_FONT, alignment=RIGHT)
        for ci in range(1, 8):
            ws.cell(row=row, column=ci).fill = LABEL_FILL
        row += 1

        _set_cell(ws, row, 6, "税费（6%）", font=TOTAL_FONT, alignment=CENTER)
        _set_cell(ws, row, 7, f"{tax_cents / 100:.2f}", font=TOTAL_FONT, alignment=RIGHT)
        for ci in range(1, 8):
            ws.cell(row=row, column=ci).fill = LABEL_FILL
        row += 1

        _set_cell(ws, row, 6, "含税合计", font=TOTAL_FONT, alignment=CENTER)
        _set_cell(ws, row, 7, f"{gross_cents / 100:.2f}", font=TOTAL_FONT, alignment=RIGHT)
        for ci in range(1, 8):
            ws.cell(row=row, column=ci).fill = LABEL_FILL
        row += 1

        # ── 调整说明 ──
        adjustments = plan.get("adjustments", [])
        if adjustments:
            row += 1
            _set_cell(ws, row, 1, "说明", font=HEADER_FONT, alignment=LEFT, border=THIN_BORDER)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
            _set_cell(ws, row, 2, "；".join(adjustments), font=BODY_FONT, alignment=LEFT, border=THIN_BORDER)

        # ── 打印设置 ──
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read()
