"""
Word document exporter.

Converts Markdown design document to .docx,
properly rendering inline syntax (bold, italic, code, links)
and applying clean formatting.
"""

import re
from io import BytesIO

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn


class WordExporter:

    @staticmethod
    def export(markdown_text: str) -> BytesIO:
        doc = Document()
        WordExporter._setup_default_font(doc)
        lines = markdown_text.split("\n")
        WordExporter._render_lines(doc, lines)
        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf

    @staticmethod
    def _setup_default_font(doc: Document) -> None:
        style = doc.styles["Normal"]
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:ascii"), "Calibri")
        rFonts.set(qn("w:hAnsi"), "Calibri")
        rFonts.set(qn("w:eastAsia"), "SimSun")

        for level in range(1, 4):
            hs = doc.styles[f"Heading {level}"]
            hs.font.bold = True
            hs.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
            sizes = {1: 18, 2: 15, 3: 13}
            hs.font.size = Pt(sizes[level])
            hRpr = hs.element.get_or_add_rPr()
            hRfonts = hRpr.get_or_add_rFonts()
            hRfonts.set(qn("w:eastAsia"), "SimHei")

    @staticmethod
    def _render_lines(doc: Document, lines: list[str]) -> None:
        in_code_block = False
        code_lines: list[str] = []
        in_table = False
        table_data: list[list[str]] = []
        list_items: list[str] = []

        for line in lines:
            stripped = line.strip()

            if stripped.startswith("```"):
                if in_code_block:
                    WordExporter._add_code_block(doc, code_lines)
                    code_lines = []
                    in_code_block = False
                else:
                    in_code_block = True
                continue
            if in_code_block:
                code_lines.append(line)
                continue

            if not stripped:
                WordExporter._flush(doc, list_items, table_data)
                list_items, table_data = [], []
                in_table = False
                continue

            if stripped.startswith("### "):
                WordExporter._flush(doc, list_items, table_data)
                list_items, table_data = [], []
                in_table = False
                WordExporter._add_heading(doc, stripped[4:], 3)
                continue
            if stripped.startswith("## "):
                WordExporter._flush(doc, list_items, table_data)
                list_items, table_data = [], []
                in_table = False
                WordExporter._add_heading(doc, stripped[3:], 2)
                continue
            if stripped.startswith("# "):
                WordExporter._flush(doc, list_items, table_data)
                list_items, table_data = [], []
                in_table = False
                WordExporter._add_heading(doc, stripped[2:], 1)
                continue

            if stripped.startswith("- ") or stripped.startswith("* "):
                if in_table and table_data:
                    WordExporter._add_table(doc, table_data)
                    table_data = []
                    in_table = False
                list_items.append(stripped[2:])
                continue

            if "|" in stripped and stripped.startswith("|"):
                if list_items:
                    WordExporter._add_list(doc, list_items)
                    list_items = []
                cells = [c.strip() for c in stripped.split("|")[1:-1]]
                if re.match(r"^[\s\-:|]+$", stripped):
                    continue
                in_table = True
                table_data.append(cells)
                continue

            WordExporter._flush(doc, list_items, table_data)
            list_items, table_data = [], []
            in_table = False
            WordExporter._add_paragraph(doc, stripped)

        if code_lines:
            WordExporter._add_code_block(doc, code_lines)
        if list_items:
            WordExporter._add_list(doc, list_items)
        if table_data:
            WordExporter._add_table(doc, table_data)

    @staticmethod
    def _flush(doc, list_items, table_data):
        if list_items:
            WordExporter._add_list(doc, list_items)
        if table_data:
            WordExporter._add_table(doc, table_data)

    # ---- inline markdown parser ----

    _INLINE_PATTERN = re.compile(
        r"(\*\*(.+?)\*\*)"          # **bold**
        r"|(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)"  # *italic* (single asterisk)
        r"|(`)(.+?)`"               # `inline code`
        r"|\[(.+?)\]\(.+?\)"        # [text](url)
    )

    @staticmethod
    def _parse_inline(text: str) -> list[tuple[str, str]]:
        """
        Parse inline markdown syntax and return list of (text, style) pairs.
        style: 'normal', 'bold', 'italic', 'code'
        """
        result: list[tuple[str, str]] = []
        pos = 0
        for m in WordExporter._INLINE_PATTERN.finditer(text):
            start = m.start()
            if start > pos:
                result.append((text[pos:start], "normal"))
            if m.group(1):  # **bold**
                result.append((m.group(2), "bold"))
            elif m.group(3):  # *italic*
                result.append((m.group(3), "italic"))
            elif m.group(5):  # `code`
                result.append((m.group(5), "code"))
            elif m.group(6):  # [text](url)
                result.append((m.group(6), "normal"))
            pos = m.end()
        if pos < len(text):
            result.append((text[pos:], "normal"))
        return result

    @staticmethod
    def _write_inline_paragraph(doc, text: str, style_name: str | None = None) -> None:
        """Add a paragraph with inline markdown rendering."""
        segments = WordExporter._parse_inline(text)
        if style_name:
            p = doc.add_paragraph(style=style_name)
        else:
            p = doc.add_paragraph()
        for seg_text, seg_style in segments:
            run = p.add_run(seg_text)
            if seg_style == "bold":
                run.bold = True
            elif seg_style == "italic":
                run.italic = True
            elif seg_style == "code":
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # ---- renderers ----

    @staticmethod
    def _add_heading(doc: Document, text: str, level: int) -> None:
        """Add heading with inline markdown stripped."""
        clean = WordExporter._strip_inline(text)
        doc.add_heading(clean, level=level)

    @staticmethod
    def _add_paragraph(doc: Document, text: str) -> None:
        WordExporter._write_inline_paragraph(doc, text)

    @staticmethod
    def _add_code_block(doc: Document, lines: list[str]) -> None:
        for line in lines:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.left_indent = Pt(12)
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    @staticmethod
    def _add_list(doc: Document, items: list[str]) -> None:
        for item in items:
            WordExporter._write_inline_paragraph(doc, item, "List Bullet")

    @staticmethod
    def _add_table(doc: Document, data: list[list[str]]) -> None:
        if not data:
            return
        rows = len(data)
        cols = max(len(r) for r in data) if data else 0
        if cols == 0:
            return
        table = doc.add_table(rows=rows, cols=cols, style="Table Grid")
        table.autofit = True
        for i, row_data in enumerate(data):
            for j, cell_text in enumerate(row_data):
                if j >= cols:
                    break
                clean = WordExporter._strip_inline(cell_text)
                table.rows[i].cells[j].text = clean
                if i == 0:
                    for paragraph in table.rows[i].cells[j].paragraphs:
                        for run in paragraph.runs:
                            run.bold = True

    @staticmethod
    def _strip_inline(text: str) -> str:
        """Remove all inline markdown markers from text."""
        t = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        t = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"\1", t)
        t = re.sub(r"`(.+?)`", r"\1", t)
        t = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", t)
        return t
