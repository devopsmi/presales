"""
Excel quotation exporter.

Exports task breakdown + summary as .xlsx,
matching the format of docs/ template.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas import EstimateResult


class ExcelExporter:

    HEADER_FILL = PatternFill("solid", fgColor="F5F7FA")
    FOOTER_FILL = PatternFill("solid", fgColor="F0F4FF")
    GRAND_FILL = PatternFill("solid", fgColor="1A1A2E")

    HEADER_FONT = Font(name="Arial", size=11, bold=True, color="555555")
    BODY_FONT = Font(name="Arial", size=10, color="333333")
    TITLE_FONT = Font(name="Arial", size=14, bold=True, color="1A1A2E")
    FOOTER_FONT = Font(name="Arial", size=10, bold=True, color="1A1A2E")
    GRAND_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")

    CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
    RIGHT = Alignment(horizontal="right", vertical="center")

    THIN_BORDER = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )

    COL_WIDTHS = [6, 16, 18, 14, 10, 36, 8, 8, 8, 14]

    @classmethod
    def export(cls, result: EstimateResult,
               company_name: str = "",
               project_name: str = "",
               quotation_unit: str = "") -> BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = "软件开发报价清单"

        ws.merge_cells("A1:J1")
        c = ws["A1"]
        c.value = "软件开发报价清单        金额：元"
        c.font = cls.TITLE_FONT
        c.alignment = cls.LEFT

        # 基本信息行（对标模板第2~3行）
        if company_name or project_name or quotation_unit:
            info_parts = []
            if company_name:
                info_parts.append(f"客户名称：{company_name}")
            if project_name:
                info_parts.append(f"项目名称：{project_name}")
            if quotation_unit:
                info_parts.append(f"报价单位：{quotation_unit}")
            ws.merge_cells("A2:J2")
            info_cell = ws["A2"]
            info_cell.value = "    ".join(info_parts)
            info_cell.font = Font(name="Arial", size=10, color="666666")
            info_cell.alignment = cls.LEFT

        for i, w in enumerate(cls.COL_WIDTHS, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        headers = ["序号", "模块", "子模块", "功能", "加工方", "功能描述", "后端", "前端", "测试", "备注"]
        header_row = 4 if (company_name or project_name or quotation_unit) else 3
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col, value=h)
            cell.font = cls.HEADER_FONT
            cell.fill = cls.HEADER_FILL
            cell.alignment = cls.CENTER
            cell.border = cls.THIN_BORDER

        row = header_row + 1

        # 按模块分组
        mod_groups: dict[str, dict[str, list]] = {}
        for t in result.tasks:
            mod_groups.setdefault(t.module, {})
            mod_groups[t.module].setdefault(t.sub_module, []).append(t)

        for mod, sub_dict in mod_groups.items():
            mod_start = row
            for sub, sub_tasks in sub_dict.items():
                sub_start = row
                for task in sub_tasks:
                    cls._write_task_row(ws, row, task)
                    row += 1
                sub_end = row - 1
                # 同 sub_module 内合并子模块列（C列）
                if sub_end > sub_start:
                    ws.merge_cells(start_row=sub_start, start_column=3, end_row=sub_end, end_column=3)
            mod_end = row - 1
            # 同 module 内合并模块列（B列）
            if mod_end > mod_start:
                ws.merge_cells(start_row=mod_start, start_column=2, end_row=mod_end, end_column=2)

        # 底部汇总区
        row += 1
        s = result.summary
        for label, b, f, t_val in [
            ("实时合计", s.total_backend_days, s.total_frontend_days, s.total_test_days),
            ("实时单价", s.backend_daily_rate, s.frontend_daily_rate, s.test_daily_rate),
            ("实时小计", s.backend_cost, s.frontend_cost, s.test_cost),
        ]:
            cls._write_footer_triple(ws, row, label, b, f, t_val)
            row += 1

        cls._write_single_value_row(ws, row, "实时费合计", s.total_labor_cost, highlight=True)
        row += 1
        cls._write_single_value_row(ws, row, f"管理费{int(s.management_fee_rate * 100)}%", s.management_fee)
        row += 1
        cls._write_single_value_row(ws, row, f"税金{int(s.tax_rate * 100)}%", s.tax)
        row += 1
        cls._write_grand_row(ws, row, "合计", s.grand_total)

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    @classmethod
    def _apply_style(cls, ws, row: int, fill=None):
        for col in range(1, 11):
            c = ws.cell(row=row, column=col)
            if fill:
                c.fill = fill
            c.border = cls.THIN_BORDER

    @classmethod
    def _write_task_row(cls, ws, row: int, task) -> None:
        values = [
            task.serial, task.module, task.sub_module, task.feature,
            task.processor, task.description,
            task.backend_days or "-", task.frontend_days or "-",
            task.test_days or "-", task.remark or "",
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = cls.BODY_FONT
            cell.border = cls.THIN_BORDER
            cell.alignment = cls.CENTER if col in (1, 7, 8, 9) else cls.LEFT

    @classmethod
    def _write_footer_triple(cls, ws, row: int, label: str, b, f, t) -> None:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cls._apply_style(ws, row)
        c = ws.cell(row=row, column=1, value=label)
        c.font = cls.FOOTER_FONT
        c.alignment = cls.LEFT
        for col, val in [(7, b), (8, f), (9, t)]:
            cell = ws.cell(row=row, column=col, value=val)
            cell.alignment = cls.CENTER
            cell.font = cls.FOOTER_FONT

    @classmethod
    def _write_single_value_row(cls, ws, row: int, label: str, value: float, highlight: bool = False) -> None:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        fill = cls.FOOTER_FILL if highlight else None
        cls._apply_style(ws, row, fill=fill)
        ws.cell(row=row, column=1, value=label).font = cls.FOOTER_FONT
        c = ws.cell(row=row, column=10, value=value)
        c.font = cls.FOOTER_FONT
        c.alignment = cls.RIGHT
        c.number_format = '#,##0'

    @classmethod
    def _write_grand_row(cls, ws, row: int, label: str, value: float) -> None:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        cls._apply_style(ws, row, fill=cls.GRAND_FILL)
        ws.cell(row=row, column=1, value=label).font = cls.GRAND_FONT
        c = ws.cell(row=row, column=10, value=value)
        c.font = cls.GRAND_FONT
        c.alignment = cls.RIGHT
        c.number_format = '#,##0'
